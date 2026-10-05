"""Local synthetic documents; no network or user documents are read."""
from pathlib import Path
from reportlab.pdfgen import canvas
from reportlab.lib.utils import ImageReader
from PIL import Image, ImageDraw
import io, json, sys

out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)
c = canvas.Canvas(str(out / 'complex.pdf'), pagesize=(612, 792))
c.setFillColorRGB(.96, .97, .99); c.rect(0, 0, 612, 792, fill=1, stroke=0)
c.setFillColorRGB(.08, .18, .32); c.rect(30, 680, 552, 80, fill=1, stroke=0)
c.setFillColorRGB(1, 1, 1); c.setFont('Helvetica-Bold', 24); c.drawString(48, 712, 'Working memory and learning')
c.beginForm('inner', 0, 0, 450, 150)
c.setFillColorRGB(.12, .42, .65); c.rect(0, 0, 450, 150, fill=1, stroke=0)
c.setFillColorRGB(1, 1, 1); c.setFont('Helvetica', 16)
c.drawString(15, 112, 'A nested Form contains this sentence.')
c.drawString(15, 80, 'Context improves long-term recall.')
c.endForm()
c.beginForm('outer', 0, 0, 500, 200); c.saveState(); c.translate(20, 20); c.doForm('inner'); c.restoreState(); c.endForm()
c.saveState(); c.translate(35, 445); c.doForm('outer'); c.restoreState()
c.setFillColorRGB(.18, .62, .41); c.circle(150, 320, 70, stroke=0, fill=1)
c.saveState(); c.setFillAlpha(.45); c.setFillColorRGB(.95, .55, .2); c.circle(220, 320, 70, stroke=0, fill=1); c.restoreState()
c.saveState(); clip=c.beginPath();clip.rect(325,285,90,90);c.clipPath(clip,stroke=0,fill=0);c.setFillColorRGB(.75,.28,.36);c.circle(330,330,75,stroke=0,fill=1);c.restoreState()
c.setFillColorRGB(.12, .16, .22); c.setFont('Helvetica', 12); c.drawString(65, 214, 'This figure shows overlapping concepts.')
c.saveState(); c.translate(525, 300); c.rotate(90); c.drawString(0, 0, 'Rotated axis label'); c.restoreState()
c.setLineWidth(1); c.setStrokeColorRGB(.35, .4, .45)
for y in [170, 140, 110]: c.line(50, y, 490, y)
for x in [50, 270, 490]: c.line(x, 110, x, 170)
c.drawString(60, 150, 'Strategy'); c.drawString(280, 150, 'Effect')
c.drawString(60, 120, 'Practice'); c.drawString(280, 120, 'Improves recall')
c.showPage()
c.setFillColorRGB(1, 1, 1); c.rect(0, 0, 612, 792, fill=1, stroke=0)
c.saveState(); c.translate(55, 500); c.scale(.9, .9); c.doForm('outer'); c.restoreState()
c.setFillColorRGB(.1, .1, .1); c.setFont('Helvetica', 18); c.drawString(50, 720, 'Shared form on a different source page')
c.showPage()
scan = Image.new('RGB', (1224, 1584), '#fbf4dd'); draw = ImageDraw.Draw(scan)
from PIL import ImageFont
font = ImageFont.truetype('/System/Library/Fonts/Supplemental/Arial.ttf', 38)
draw.text((95, 100), 'Scanned lesson notes', font=font, fill='#153d52')
draw.text((95, 200), 'Spaced practice improves retention.', font=font, fill='#153d52')
draw.rectangle((95, 360, 600, 780), fill='#409d82'); draw.ellipse((650, 400, 1120, 840), fill='#edbc49')
draw.text((95, 1000), 'Figures are preserved without analysis.', font=font, fill='#153d52')
scan.save(out / 'scan.png'); c.drawImage(ImageReader(scan), 0, 0, 612, 792); c.showPage(); c.save()
from pypdf import PdfReader, PdfWriter
rotated=PdfWriter();rotated.add_page(PdfReader(str(out/'complex.pdf')).pages[0]);rotated.pages[0].rotate(90)
with (out/'rotated.pdf').open('wb') as file: rotated.write(file)
from reportlab.lib.pdfencrypt import StandardEncryption
encrypted = canvas.Canvas(str(out / 'encrypted.pdf'), pagesize=(612, 792), encrypt=StandardEncryption('fixture-password',canPrint=0))
encrypted.drawString(50,700,'Protected lesson');encrypted.save()
(out / 'corrupt.pptx').write_bytes(b'corrupt slide fixture')
(out / 'lesson.md').write_text('# Working memory\n\nContext matters. **Practice** supports recall.\n\n| Strategy | Effect |\n| --- | --- |\n| Practice | Retention |\n\n```python\nprint("preserve code")\n```\n\n![Learning illustration](scan.png)\n', encoding='utf-8')
(out / 'lesson.txt').write_text('Working memory\n\nContext improves learning.\nPractice supports recall.\n', encoding='utf-8')
package = out / 'lesson.ulnote'; (package / 'assets').mkdir(parents=True, exist_ok=True)
(package / 'assets' / 'scan.png').write_bytes((out / 'scan.png').read_bytes())
def block(kind, text='', **extras):
    import uuid
    value = dict(id=str(uuid.uuid5(uuid.NAMESPACE_URL, f'ulecture-fixture-{kind}')),kind=kind,text=text,level=2,checked=False,ordered=False,cells=[],imageWidth=480,codeLanguage='',links=[])
    value.update(extras); return value
note = dict(format='ulecture-block-note',formatVersion=1,id='fixture-note',title='Learning notes',revision=1,savedAt=0,blocks=[block('heading','Working memory'),block('paragraph','Context improves learning.'),block('table',cells=[['Strategy','Effect'],['Practice','Retention']]),block('code','print("preserve code")',codeLanguage='python'),block('image','Learning illustration',resource='assets/scan.png')])
(package / 'note.json').write_text(json.dumps(note), encoding='utf-8')
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.oxml.xmlchemy import OxmlElement
    p = Presentation(); p.slide_width = Inches(10); p.slide_height = Inches(7.5)
    s = p.slides.add_slide(p.slide_layouts[6]); title = s.shapes.add_textbox(Inches(.5), Inches(.4), Inches(9), Inches(1))
    title.text = 'Working memory / 日本語 / 中文'; title.text_frame.paragraphs[0].font.size = Pt(28)
    shape = s.shapes.add_textbox(Inches(.5), Inches(1.6), Inches(9), Inches(.7)); shape.text = 'Context improves learning. Practice supports recall.'
    shape.text_frame.paragraphs[0].font.name = 'FixtureMissingFont'
    table = s.shapes.add_table(2, 2, Inches(.5), Inches(2.8), Inches(5), Inches(1)).table
    for (r, col, value) in [(0,0,'Strategy'),(0,1,'Effect'),(1,0,'Practice'),(1,1,'Retention')]: table.cell(r,col).text = value
    s.shapes.add_picture(str(out / 'scan.png'), Inches(6), Inches(3), width=Inches(2), height=Inches(2.5))
    s.notes_slide.notes_text_frame.text = 'Speaker notes are separate from slide content.'
    p.save(out / 'lesson.pptx')
    # A separate fidelity fixture adds formula/chart/animation/transition and
    # a rotated shape without changing the baseline single-slide contract.
    formula = s.shapes.add_textbox(Inches(.5), Inches(5.9), Inches(3), Inches(.5)); formula.text = 'E = mc²'
    chart_data = CategoryChartData(); chart_data.categories=['Practice','Rest']; chart_data.add_series('Recall',[8,4])
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(3.8), Inches(5.2), Inches(3.2), Inches(1.7), chart_data)
    rotated = s.shapes.add_textbox(Inches(8.5), Inches(5), Inches(1), Inches(1)); rotated.text='Axis'; rotated.rotation=90
    from lxml import etree
    ns='http://schemas.openxmlformats.org/presentationml/2006/main'
    transition=etree.fromstring(f'<p:transition xmlns:p="{ns}" spd="slow"><p:fade/></p:transition>');s._element.append(transition)
    timing=etree.fromstring(f'<p:timing xmlns:p="{ns}"><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst><p:par><p:cTn id="2" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst><p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="3" dur="1000"/><p:tgtEl><p:spTgt spid="{title.shape_id}"/></p:tgtEl></p:cBhvr></p:animEffect></p:childTnLst></p:cTn></p:par></p:childTnLst></p:cTn></p:par></p:tnLst></p:timing>');s._element.append(timing)
    p.save(out / 'fidelity.pptx')
except ImportError:
    pass
print(json.dumps({'fixtureDirectory':str(out), 'files':[p.name for p in out.iterdir()]}))
